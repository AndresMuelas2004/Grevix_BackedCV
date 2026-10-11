import io

import docx
from docx.shared import Inches, Pt

from backedcv.domain.models import CVDocument, DocContact, DocItem, DocLine, DocSection
from backedcv.rendering import render_docx, render_pdf, roundtrip_check

EXPECTED_LINES = [
    "José Muñoz",
    "jose.munoz@example.com | +34 600 123 456 | Málaga, Spain",
    "Summary",
    "R&D engineer with C++ and C#",
    "Experience",
    "Data Engineer, Acme Corp, Jan 2020 – Present",
    "Built ETL pipelines in C++ and Python",
    "Led an R&D team of 4 people",
    "Developer, Beta Labs, 2017 – 2019",
    "Maintained C# reports for Niño and Peña clients",
    "Skills",
    "C++, C#, R&D, SQL, Python",
]


def _line(text: str, kind: str = "text") -> DocLine:
    return DocLine(text=text, origin="L001", kind=kind)


def _cv() -> CVDocument:
    return CVDocument(
        contact=DocContact(
            name="José Muñoz",
            details=["jose.munoz@example.com", "+34 600 123 456", "Málaga, Spain"],
            origins=["L001", "L002"],
        ),
        sections=[
            DocSection(
                id="summary",
                type="summary",
                heading="Summary",
                items=[DocItem(id="summary-i1", lines=[_line("R&D engineer with C++ and C#")])],
            ),
            DocSection(
                id="experience",
                type="experience",
                heading="Experience",
                items=[
                    DocItem(
                        id="job-1",
                        lines=[
                            _line("Data Engineer, Acme Corp, Jan 2020 – Present"),
                            _line("Built ETL pipelines in C++ and Python", "bullet"),
                            _line("Led an R&D team of 4 people", "bullet"),
                        ],
                    ),
                    DocItem(
                        id="job-2",
                        lines=[
                            _line("Developer, Beta Labs, 2017 – 2019"),
                            _line("Maintained C# reports for Niño and Peña clients", "bullet"),
                        ],
                    ),
                ],
            ),
            DocSection(
                id="skills",
                type="skills",
                heading="Skills",
                items=[DocItem(id="skills-i1", lines=[_line("C++, C#, R&D, SQL, Python")])],
            ),
        ],
    )


def test_docx_reopens_with_every_line_in_order():
    document = docx.Document(io.BytesIO(render_docx(_cv())))

    assert [p.text for p in document.paragraphs if p.text] == EXPECTED_LINES


def test_docx_is_plain_one_column_with_no_tables_images_or_boxes():
    document = docx.Document(io.BytesIO(render_docx(_cv())))
    section = document.sections[0]
    paragraphs = {p.text: p for p in document.paragraphs}

    assert document.tables == []
    assert len(document.inline_shapes) == 0
    assert not document.element.body.xpath(".//w:txbxContent")
    assert section.header.is_linked_to_previous and section.footer.is_linked_to_previous
    assert paragraphs["Built ETL pipelines in C++ and Python"].style.name == "List Bullet"
    assert paragraphs["Experience"].runs[0].bold
    assert document.styles["Normal"].font.name == "Calibri"
    assert document.styles["Normal"].font.size == Pt(10.5)
    assert section.left_margin == Inches(0.7)


def test_pdf_reads_back_with_every_line():
    pdf = render_pdf(_cv())

    assert pdf.startswith(b"%PDF")
    ok, issues = roundtrip_check(pdf, _cv())
    assert ok, issues


def test_roundtrip_reports_a_line_missing_from_the_pdf():
    pdf = render_pdf(_cv())
    doc = _cv()
    doc.sections[0].items[0].lines.append(_line("This line was never rendered"))

    ok, issues = roundtrip_check(pdf, doc)

    assert not ok
    assert issues == ["This line was never rendered"]
