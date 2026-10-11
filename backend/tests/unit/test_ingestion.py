import io
from pathlib import Path

import docx
import pytest
import typst
from docx.oxml import parse_xml

from backedcv.ingestion import IngestionError, read_cv, read_cv_text

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
PAGE = "#set page(width: 210mm, height: 297mm, margin: 20mm)\n"
NARROW_PAGE = "#set page(width: 90mm, height: 200mm, margin: 10mm)\n"
BULLET = (
    "Built ETL jobs in Python that load two million rows a day into the PostgreSQL warehouse "
    "for the analytics team"
)
CURATED_DIR = Path(__file__).resolve().parents[2] / "samples" / "curated"


def _pdf(tmp_path, markup: str) -> bytes:
    source = tmp_path / "cv.typ"
    source.write_text(markup, encoding="utf-8")
    return typst.compile(str(source))


def test_pasted_text_gets_numbered_lines_without_bullets():
    result = read_cv_text("Jane Doe\n\n• Built APIs with FastAPI\n- Led a team of 4\n")

    assert [(line.id, line.text) for line in result.lines] == [
        ("L001", "Jane Doe"),
        ("L002", "Built APIs with FastAPI"),
        ("L003", "Led a team of 4"),
    ]


def test_pasted_text_drops_control_characters():
    result = read_cv_text("Jane\x00\tDoe\n\x1b\x07\n\x1bBuilt APIs with FastAPI\x7f\n")

    assert [(line.id, line.text) for line in result.lines] == [
        ("L001", "Jane Doe"),
        ("L002", "Built APIs with FastAPI"),
    ]


def test_pasted_text_drops_unpaired_surrogates():
    result = read_cv_text("Jane\ud800 Doe 😀\n\udc00\n")

    assert [line.text for line in result.lines] == ["Jane Doe \U0001f600"]
    assert result.size == len("Jane Doe \U0001f600".encode())  # the size of the canonical text


def test_pdf_lines_in_reading_order_with_bold_headings(tmp_path):
    data = _pdf(
        tmp_path,
        "#set page(width: 210mm, height: 297mm)\n"
        "*Jane Doe*\n\n"
        "jane\\@example.com\n\n"
        "*Experience*\n\n"
        "Data Engineer, Acme Corp, Jan 2020 - Present\n\n"
        "- Built ETL pipelines in Python and SQL\n",
    )

    result = read_cv(data, "cv.pdf")
    texts = [line.text for line in result.lines]

    assert result.kind == "pdf"
    assert texts[0] == "Jane Doe"
    assert "Data Engineer, Acme Corp, Jan 2020 - Present" in texts
    assert texts[-1] == "Built ETL pipelines in Python and SQL"
    assert next(line for line in result.lines if line.text == "Experience").bold


def _pdf_texts(tmp_path, markup: str) -> list[str]:
    return [line.text for line in read_cv(_pdf(tmp_path, markup), "cv.pdf").lines]


def test_wrapped_bullet_becomes_one_line(tmp_path):
    assert _pdf_texts(tmp_path, NARROW_PAGE + f"- {BULLET}\n") == [BULLET]


def test_consecutive_bullets_stay_separate_lines(tmp_path):
    second = "Led a team of four analysts through the warehouse migration to the cloud platform"

    texts = _pdf_texts(tmp_path, NARROW_PAGE + f"- {BULLET}\n- {second}\n")

    assert texts == [BULLET, second]


def test_literal_bullet_wrapped_to_the_margin_stays_whole(tmp_path):
    # A "•" typed as text: its wrapped line starts at the margin, not under the bullet's text.
    markup = NARROW_PAGE + (
        "• Containerized three internal services with Docker and exposed them through FastAPI.\n"
    )

    assert _pdf_texts(tmp_path, markup) == [
        "Containerized three internal services with Docker and exposed them through FastAPI."
    ]


def test_wrapped_paragraph_becomes_one_line(tmp_path):
    summary = (
        "Software engineer with three years of data and backend work, now moving into AI "
        "engineering. I build LLM features with agent frameworks and tool calling."
    )

    assert _pdf_texts(tmp_path, NARROW_PAGE + f"{summary}\n") == [summary]


def test_one_skill_per_line_stays_separate(tmp_path):
    # The bullet sets the column's right edge; the short skill lines stop far short of it.
    markup = NARROW_PAGE + f"- {BULLET}\n\nPython \\\nSQL \\\nDocker\n"

    assert _pdf_texts(tmp_path, markup) == [BULLET, "Python", "SQL", "Docker"]


def test_heading_right_after_a_paragraph_is_not_merged(tmp_path):
    markup = (
        PAGE + "We build data pipelines for the analytics team and the finance group \\\n"
        "*Experience*\n"
    )

    assert _pdf_texts(tmp_path, markup) == [
        "We build data pipelines for the analytics team and the finance group",
        "Experience",
    ]


def test_date_line_is_not_merged_into_the_title_above(tmp_path):
    markup = PAGE + "Data Engineer, Acme Corp \\\nJan 2020 - Present\n"

    assert _pdf_texts(tmp_path, markup) == ["Data Engineer, Acme Corp", "Jan 2020 - Present"]


def test_hyphen_at_a_line_break_joins_a_lowercase_word(tmp_path):
    # The first line is the widest in its block, so it reaches the edge like a real wrap.
    markup = PAGE + (
        "Built ETL pipelines for the analytics team across every region and the data-\\\n"
        "driven dashboard for the team\n"
    )

    assert _pdf_texts(tmp_path, markup) == [
        "Built ETL pipelines for the analytics team across every region and the data-driven "
        "dashboard for the team"
    ]


def test_bold_title_wrapped_after_a_comma_stays_one_line(tmp_path):
    # A bold degree title too wide for a narrow column: the first line ends with a comma.
    markup = "#set page(width: 78mm, height: 200mm, margin: 10mm)\n" + (
        "*BSc Computer Science, Universidad Example, Valencia*\n"
    )

    assert _pdf_texts(tmp_path, markup) == ["BSc Computer Science, Universidad Example, Valencia"]


def test_curated_sample_keeps_bullets_and_wrapped_paragraphs_whole():
    texts = [line.text for line in read_cv((CURATED_DIR / "cv.pdf").read_bytes(), "cv.pdf").lines]
    source = (CURATED_DIR / "cv.txt").read_text(encoding="utf-8").splitlines()
    bullets = [line[2:] for line in source if line.startswith("- ")]
    paragraphs = [
        "Software engineer with three years of data and backend work, now moving into AI "
        "engineering. I build LLM features with agent frameworks and tool calling, measure them "
        "on labeled test sets, and serve them through FastAPI.",
        "Python, SQL, PostgreSQL, Pandas, FastAPI, Docker, Git, OpenAI API, pgvector, "
        "prompt evaluation",
    ]

    assert [text for text in bullets + paragraphs if text not in texts] == []
    assert "Wrote ETL jobs in Python that load 2 million rows a day into PostgreSQL." in texts
    assert not {"PostgreSQL.", "analysts.", "FastAPI."} & set(texts)
    assert {"Valencia, Spain", "nuria.calvo@example.com"} <= set(texts)


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_reads_tables_headers_and_content_controls():
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe | jane@example.com"
    document.add_heading("Experience", level=1)
    document.add_paragraph("Data Engineer, Acme Corp")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    sdt = parse_xml(
        f"<w:sdt {W_NS}><w:sdtContent><w:p><w:r><w:t>Skills inside a content control"
        "</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    document.element.body.insert(len(document.element.body) - 1, sdt)

    result = read_cv(_docx_bytes(document), "cv.docx")
    by_text = {line.text: line for line in result.lines}

    assert by_text["Jane Doe | jane@example.com"].source == "header"
    assert by_text["Experience"].bold
    assert by_text["Python"].source == "table"
    assert "Skills inside a content control" in by_text


def test_docx_without_a_body_is_unreadable():
    document = docx.Document()
    document.element.remove(document.element.body)

    with pytest.raises(IngestionError) as error:
        read_cv(_docx_bytes(document), "cv.docx")
    assert error.value.code == "cv_unreadable"
    assert error.value.status == 422


def test_rejects_unsupported_files():
    with pytest.raises(IngestionError) as error:
        read_cv(b"just some bytes", "cv.txt")
    assert error.value.code == "cv_unsupported_type"


def test_image_only_pdf_asks_for_pasted_text(tmp_path):
    data = _pdf(tmp_path, "#rect(width: 4cm, height: 2cm, fill: gray)")

    with pytest.raises(IngestionError) as error:
        read_cv(data, "scan.pdf")
    assert error.value.code == "cv_image_only"
