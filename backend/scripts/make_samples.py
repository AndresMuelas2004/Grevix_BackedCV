"""Render the curated sample CV as a two-column PDF (Typst) and a one-column DOCX (python-docx).

Reads backend/samples/curated/cv.txt and writes cv.pdf and cv.docx next to it.
Run from backend/: uv run python scripts/make_samples.py
"""

import re
import tempfile
from pathlib import Path

import docx
import typst
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CURATED = Path(__file__).resolve().parents[1] / "samples" / "curated"
SOURCE = CURATED / "cv.txt"
PDF_OUT = CURATED / "cv.pdf"
DOCX_OUT = CURATED / "cv.docx"

LABELS = {
    "CONTACT": "Contact",
    "SUMMARY": "Summary",
    "EXPERIENCE": "Experience",
    "PROJECTS": "Projects",
    "EDUCATION": "Education",
    "SKILLS": "Skills",
    "LANGUAGES": "Languages",
}
# PDF: a narrow left column and a wide right column, under a full-width header.
LEFT_COLUMN = ("CONTACT", "EDUCATION", "SKILLS", "LANGUAGES")
RIGHT_COLUMN = ("SUMMARY", "EXPERIENCE", "PROJECTS")
# In these sections the first plain line of an entry is its title (job, project, degree).
ENTRY_SECTIONS = ("EXPERIENCE", "PROJECTS", "EDUCATION")
DATE_LINE = re.compile(r"^[A-Z][a-z]{2} \d{4} - (?:[A-Z][a-z]{2} \d{4}|Present)$")
TYPST_CALL = {"title": "entry-title", "date": "entry-date", "bullet": "bullet", "text": "para"}

TYPST_TEMPLATE = """\
#set document(title: {title}, author: {name})
#set page(paper: "a4", margin: (x: 1.8cm, y: 1.6cm))
#set text(font: "Libertinus Serif", size: 10.5pt, ligatures: false, hyphenate: false, lang: "en")
#set smartquote(enabled: false)
#set par(leading: 0.65em)

#let section-title(s) = block(above: 1.2em, below: 0.6em, width: 100%)[
  #text(size: 11pt, weight: "bold", s)
  #v(-0.2em)
  #line(length: 100%, stroke: 0.6pt + luma(120))
]
#let entry-title(s) = block(above: 0.9em, below: 0.15em)[#text(weight: "bold", s)]
#let entry-date(s) = block(below: 0.25em)[#text(size: 9pt, fill: luma(90), s)]
#let bullet(s) = block(below: 0.35em)[#grid(columns: (0.9em, 1fr), [•], [#s])]
#let para(s) = block(below: 0.6em)[#s]

#text(size: 24pt, weight: "bold", {name})
#v(0.1em)
#text(size: 11.5pt, fill: luma(60), {headline})
#v(0.5em)
#line(length: 100%, stroke: 1pt)
#v(0.4em)

#grid(
  columns: (30%, 1fr),
  column-gutter: 0.8cm,
  [
{left}
  ],
  [
{right}
  ],
)
"""


def parse(text: str) -> tuple[str, str, list[tuple[str, list[str]]]]:
    """Return the name, the headline and the (heading, lines) sections in source order."""
    header: list[str] = []
    sections: list[tuple[str, list[str]]] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line in LABELS:
            sections.append((line, []))
        elif sections:
            sections[-1][1].append(line)
        else:
            header.append(line)
    if len(header) != 2:
        raise ValueError(f"{SOURCE.name}: expected a name and a headline before the first heading")
    return header[0], header[1], sections


def blocks(heading: str, lines: list[str]) -> list[tuple[str, str]]:
    """Classify each line as a bullet, a date range, an entry title or plain text."""
    result = []
    for line in lines:
        if line.startswith("- "):
            result.append(("bullet", line[2:]))
        elif DATE_LINE.match(line):
            result.append(("date", line))
        elif heading in ENTRY_SECTIONS:
            result.append(("title", line))
        else:
            result.append(("text", line))
    return result


def typst_string(text: str) -> str:
    """A Typst string literal: the text shows as plain words and is never read as markup."""
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'


def typst_column(headings: tuple[str, ...], by_heading: dict[str, list[str]]) -> str:
    lines = []
    for heading in headings:
        if heading not in by_heading:
            continue
        lines.append(f"    #section-title({typst_string(LABELS[heading])})")
        for kind, text in blocks(heading, by_heading[heading]):
            lines.append(f"    #{TYPST_CALL[kind]}({typst_string(text)})")
    return "\n".join(lines)


def build_typst(name: str, headline: str, sections: list[tuple[str, list[str]]]) -> str:
    by_heading = dict(sections)
    return TYPST_TEMPLATE.format(
        title=typst_string(f"{name} - CV"),
        name=typst_string(name),
        headline=typst_string(headline),
        left=typst_column(LEFT_COLUMN, by_heading),
        right=typst_column(RIGHT_COLUMN, by_heading),
    )


def _plain_font(style, font_name: str) -> None:
    """Set a font by name and drop the theme references that the default template adds."""
    style.font.name = font_name
    fonts = style.element.rPr.rFonts
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)


def render_docx(name: str, headline: str, sections: list[tuple[str, list[str]]]) -> None:
    document = docx.Document()
    document.core_properties.title = f"{name} - CV"
    document.core_properties.author = name
    page = document.sections[0]
    page.left_margin = page.right_margin = Cm(2.0)
    page.top_margin = page.bottom_margin = Cm(1.6)

    normal = document.styles["Normal"]
    _plain_font(normal, "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)

    heading_style = document.styles["Heading 1"]
    _plain_font(heading_style, "Arial")
    heading_style.font.size = Pt(11)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.paragraph_format.space_before = Pt(10)
    heading_style.paragraph_format.space_after = Pt(3)

    name_run = document.add_paragraph().add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(18)
    document.add_paragraph(headline)

    for heading, lines in sections:
        document.add_heading(LABELS[heading], level=1)
        if heading == "CONTACT":
            document.add_paragraph(" | ".join(lines))
            continue
        for kind, text in blocks(heading, lines):
            if kind == "title":
                paragraph = document.add_paragraph()
                paragraph.add_run(text).bold = True
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(0)
            elif kind == "date":
                paragraph = document.add_paragraph(text)
                paragraph.paragraph_format.space_after = Pt(1)
            elif kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            else:
                document.add_paragraph(text)
    document.save(DOCX_OUT)


def main() -> None:
    name, headline, sections = parse(SOURCE.read_text(encoding="utf-8"))
    markup = build_typst(name, headline, sections)
    with tempfile.TemporaryDirectory() as workdir:
        source = Path(workdir) / "cv.typ"
        source.write_text(markup, encoding="utf-8")
        # Typst's embedded fonts only: system fonts are ignored.
        PDF_OUT.write_bytes(typst.compile(str(source), ignore_system_fonts=True))
    render_docx(name, headline, sections)
    print(f"wrote {PDF_OUT.name} and {DOCX_OUT.name} in {CURATED}")


if __name__ == "__main__":
    main()
