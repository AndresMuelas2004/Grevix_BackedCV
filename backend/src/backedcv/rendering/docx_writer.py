"""Render the rebuilt CV to DOCX: one column, plain paragraphs, no tables, images or headers."""

import io
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Inches, Mm, Pt

from backedcv.domain.models import CVDocument

FONT_NAME = "Calibri"
BODY_SIZE = Pt(10.5)
NAME_SIZE = Pt(14)
MARGIN = Inches(0.7)
CONTACT_SEPARATOR = " · "


def render_docx(doc: CVDocument) -> bytes:
    document = Document()
    _page_and_styles(document)
    _metadata(document, doc)

    if doc.contact.name:
        _name(document, doc.contact.name)
    if doc.contact.headline:
        document.add_paragraph(doc.contact.headline)
    details = [detail for detail in doc.contact.details if detail.strip()]
    if details:
        document.add_paragraph(CONTACT_SEPARATOR.join(details))
    for section in doc.sections:
        if section.heading:
            _heading(document, section.heading)
        for item in section.items:
            for line in item.lines:
                style = "List Bullet" if line.kind == "bullet" else None
                document.add_paragraph(line.text, style=style)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _page_and_styles(document: WordDocument) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(2)


def _metadata(document: WordDocument, doc: CVDocument) -> None:
    props = document.core_properties
    name = doc.contact.name or ""
    props.title = name or "CV"
    props.author = name
    props.last_modified_by = name
    props.created = props.modified = datetime.now(UTC)


def _name(document: WordDocument, name: str) -> None:
    run = document.add_paragraph().add_run(name)
    run.bold = True
    run.font.size = NAME_SIZE


def _heading(document: WordDocument, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text).bold = True
